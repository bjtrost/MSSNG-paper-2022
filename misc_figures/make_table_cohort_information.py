#!/usr/bin/env python3

#####################################################
# make a table summarizing the individuals in MSSNG #
#####################################################

import argparse
import docx
import docx_lib
import pandas
from collections import defaultdict
import docx_lib

####################################
### Parse command-line arguments ###
parser = argparse.ArgumentParser(formatter_class=argparse.ArgumentDefaultsHelpFormatter)
parser.add_argument("metadata_filename", type=str)
args = parser.parse_args()
#####################################

# Testing status: checked all columns manually for one cohort

metadata = pandas.read_csv(args.metadata_filename, sep="\t", dtype=str, keep_default_na=False)


relations = ["ASD-affected child", "mother", "father", "unaffected sibling", "other"]

cohorts = metadata["Cohort"].unique()
sequencing_platforms = sorted(metadata["Platform"].unique())
columns = ["Cohort"] + relations + sequencing_platforms + ["Total ASD genomes", "Total genomes"]
output_df = pandas.DataFrame(columns=columns)

total = defaultdict(int)
total["Cohort"] = "Total"
for cohort in cohorts:
    to_insert = {}
    to_insert["Cohort"] = cohort

    for relation in relations:
        if relation == "ASD-affected child":
            to_insert[relation] = str(len(metadata.loc[(metadata["Cohort"] == cohort) & (metadata["Affection"] == "2") & ( (metadata["Relation"] == "proband") | (metadata["Relation"] == "affected sibling"))]))
        else:
            to_insert[relation] = str(len(metadata.loc[(metadata["Cohort"] == cohort) & (metadata["Relation"] == relation)]))
        total[relation] += int(to_insert[relation])

    for platform in sequencing_platforms:
        to_insert[platform] = str(len(metadata.loc[(metadata["Cohort"] == cohort) & (metadata["Platform"] == platform)]))
        total[platform] += int(to_insert[platform])

    to_insert["Total ASD genomes"] = str(len(metadata.loc[(metadata["Cohort"] == cohort) & (metadata["Affection"] == "2")]))
    to_insert["Total genomes"] = len(metadata.loc[metadata["Cohort"] == cohort])
    total["Total ASD genomes"] += int(to_insert["Total ASD genomes"])
    total["Total genomes"] += int(to_insert["Total genomes"])
    output_df = output_df.append(to_insert, ignore_index=True)

output_df.sort_values("Total genomes", inplace=True, ascending=False)
for k in total:
    if k == "Cohort": continue
    total[k] = str(int(total[k]))

output_df = output_df.append(total, ignore_index=True)
table, doc = docx_lib.df_to_docx(output_df)

for j in range(1, output_df.shape[-1]):
    docx_lib.set_vertical_cell_direction(cell=table.cell(0, j), direction="btLr")

for row in range(output_df.shape[0]):
    table.cell(row+1, 0).width = docx.shared.Inches(6)

last_row = output_df.shape[0]
for col in range(output_df.shape[-1]):
    table.cell(0, col).paragraphs[0].runs[0].font.size = docx.shared.Pt(10)
    table.cell(last_row, col).paragraphs[0].runs[0].font.bold = True

doc.save("MSSNG_DB6_cohort_information_relations.docx")
